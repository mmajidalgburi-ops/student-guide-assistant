import os
from pathlib import Path
import streamlit as st
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document
import time

st.set_page_config(page_title="المستشار الإداري", layout="wide")
st.title("🧑‍💼 المستشار الإداري - دليل الطلبة")

BASE = Path(__file__).parent

# قراءة المفتاح من إعدادات Streamlit
try:
    os.environ["GOOGLE_API_KEY"] = st.secrets["GOOGLE_API_KEY"]
except:
    st.error("❌ لم يتم العثور على المفتاح. يرجى إضافته في إعدادات Secrets.")
    st.stop()

# البحث عن ملف الوورد
docx_files = [f for f in BASE.glob("*.docx") if not f.name.startswith("~$")]
if not docx_files:
    st.error("❌ لم أجد أي ملف وورد في المجلد.")
    st.stop()
FILE = docx_files[0]

@st.cache_resource
def build():
    doc = docx.Document(str(FILE))
    text_list = []
    
    # قراءة الفقرات النصية
    for p in doc.paragraphs:
        if p.text.strip():
            text_list.append(p.text)
            
    # قراءة الجداول
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_list.append(cell.text)
                    
    text = "\n".join(text_list)
    
    splitter = RecursiveCharacterTextSplitter(chunk_size=600, chunk_overlap=150)
    chunks = splitter.split_documents([Document(page_content=text, metadata={"source": FILE.name})])
    
    emb = HuggingFaceEmbeddings(model_name="intfloat/multilingual-e5-small")
    return Chroma.from_documents(chunks, emb), len(chunks)

db, n = build()
st.caption(f"📄 {FILE.name} | 🧠 {n} مقطع جاهز")

if "messages" not in st.session_state:
    st.session_state.messages = []

for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])

def generate_answer(q):
    # 1. إعادة صياغة السؤال: تحويل اللغة البسيطة/العامية إلى مصطلحات رسمية
    rewrite_prompt = (
        "أنت خبير في الشؤون الأكاديمية. المستخدم يسأل بلغة بسيطة أو عامية.\n"
        "قم بتحويل سؤاله إلى سؤال رسمي يستخدم المصطلحات الدقيقة الموجودة في دليل الجامعات.\n"
        "مثال: إذا سأل 'الموظف يدرس'، حولها إلى 'ضوابط الجمع بين الوظيفة والدراسة'.\n"
        "مثال: إذا سأل 'أنقل من مرحلة'، حولها إلى 'ضوابط انتقال الطلبة بين المراحل'.\n"
        "أعطني السؤال الرسمي فقط بدون أي مقدمات.\n\n"
        f"سؤال المستخدم: {q}"
    )
    try:
        rw_llm = ChatGoogleGenerativeAI(model="gemini-3.6-flash", temperature=0)
        formal_q = rw_llm.invoke(rewrite_prompt).content
        if isinstance(formal_q, list): formal_q = formal_q[0].get('text', str(formal_q))
    except:
        formal_q = q

    # 2. البحث في الملف باستخدام السؤال الرسمي
    results = db.similarity_search("query: " + formal_q, k=15)
    context = "\n\n---\n\n".join([d.page_content for d in results])

    # 3. الإجابة بلغة بسيطة ودقيقة (ذكر القاعدة أولاً ثم الاستثناءات)
    prompt = (
        "أنت مستشار إداري ذكي ودقيق جداً. أجب على سؤال المستخدم بالعربية بناءً على النص المرفق.\n"
        "إذا كان السؤال يتعلق بالموظفين أو الدراسة، اذكر القاعدة العامة أولاً (هل يجوز أم لا؟)، ثم اذكر الاستثناءات إن وجدت في النص.\n"
        "اجعل إجابتك بسيطة ومباشرة (ابدأ بـ: نعم يجوز، أو: لا يجوز، ثم اشرح).\n"
        "إذا لم يوجد النص نهائياً، اعتذر بلباقة.\n\n"
        f"النص:\n{context}\n\nسؤال المستخدم: {q}"
    )
    
    MODELS = ["gemini-3.6-flash", "gemini-3.5-flash-lite", "gemini-3.7-flash"]
    for model_name in MODELS:
        try:
            llm = ChatGoogleGenerativeAI(model=model_name, temperature=0, max_retries=1)
            response = llm.invoke(prompt)
            if isinstance(response.content, list) and len(response.content) > 0:
                return response.content[0].get('text', str(response.content))
            return response.content
        except Exception:
            continue 
            
    return "عذراً، النماذج مشغولة حالياً. حاول بعد قليل."

# واجهة الدردشة
if q := st.chat_input("اكتب سؤالك هنا..."):
    st.session_state.messages.append({"role": "user", "content": q})
    with st.chat_message("user"):
        st.markdown(q)

    with st.chat_message("assistant"):
        with st.spinner("⏳ يبحث في الملف ويجيب..."):
            ans = generate_answer(q)
            st.markdown(ans)
            st.session_state.messages.append({"role": "assistant", "content": ans})
